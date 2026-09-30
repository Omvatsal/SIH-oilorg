"""FastAPI entry point for the offline Phase 1 demonstration."""

import argparse
import os
import sys
from datetime import date

from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from construction_reconciliation.demo.importer import import_csv, import_xer, schedule_summary
from construction_reconciliation.demo.pipeline import DemoStore

PRODUCT_NAME = os.getenv("PRODUCT_NAME", "SETU")
app = FastAPI(title=f"{PRODUCT_NAME} Phase 1 API")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
    allow_methods=["*"],
    allow_headers=["*"],
)
store = DemoStore()


class MessageRequest(BaseModel):
    schedule_id: str
    text: str
    area: str | None = None
    received_at: date | None = None


@app.get("/api/config")
def app_config() -> dict[str, str]:
    return {"product_name": PRODUCT_NAME}


@app.get("/health")
def health() -> dict[str, str]:
    """Return process health without depending on external services."""
    return {"status": "ok"}


@app.post("/schedule/import")
async def import_schedule(file: UploadFile):
    suffix = (file.filename or "").lower().rsplit(".", 1)[-1]
    if suffix not in {"csv", "xer"}:
        raise HTTPException(status_code=422, detail="Only CSV and XER schedules are supported")
    try:
        content = (await file.read()).decode("utf-8")
        schedule = import_xer(content) if suffix == "xer" else import_csv(content)
    except (KeyError, UnicodeDecodeError, ValueError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    store.add_schedule(schedule)
    return schedule_summary(schedule)


@app.post("/ingest/message")
def ingest_message(payload: MessageRequest):
    if payload.schedule_id not in store.schedules:
        raise HTTPException(status_code=404, detail="Schedule not found")
    return store.process(payload.schedule_id, payload.text, area=payload.area, received_at=payload.received_at)


@app.post("/ingest/file")
async def ingest_file(schedule_id: str, file: UploadFile, area: str | None = None, received_at: date | None = None):
    if schedule_id not in store.schedules:
        raise HTTPException(status_code=404, detail="Schedule not found")
    suffix = (file.filename or "").lower().rsplit(".", 1)[-1]
    raw = await file.read()
    if suffix in {"txt", "csv"}:
        text = raw.decode("utf-8")
    elif suffix == "xlsx":
        try:
            from openpyxl import load_workbook
            from io import BytesIO

            workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
            try:
                text = "\n".join(" | ".join(str(cell) for cell in row if cell is not None) for row in workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
        except Exception as exc:
            raise HTTPException(status_code=422, detail="Could not read Excel report") from exc
    else:
        raise HTTPException(status_code=422, detail="Only txt, csv, and xlsx reports are supported")
    return store.process(schedule_id, text, area=area, received_at=received_at)


@app.get("/workspace/{report_id}")
def get_workspace(report_id: str):
    try:
        return store.workspaces[report_id]
    except KeyError as exc:
        raise HTTPException(status_code=404, detail="Report not found") from exc


def _patch_action(patch_id: str, action: str):
    try:
        return store.patch_action(patch_id, action)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail="Patch not found") from exc
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


@app.post("/patches/{patch_id}/accept")
def accept_patch(patch_id: str):
    return _patch_action(patch_id, "accept")


@app.post("/patches/{patch_id}/reject")
def reject_patch(patch_id: str):
    return _patch_action(patch_id, "reject")


@app.post("/patches/{patch_id}/undo")
def undo_patch(patch_id: str):
    return _patch_action(patch_id, "undo")


@app.get("/audit/{patch_id}")
def get_audit(patch_id: str):
    if patch_id not in store.audit:
        raise HTTPException(status_code=404, detail="Patch not found")
    return {"patch_id": patch_id, "entries": store.audit[patch_id]}


@app.get("/metrics")
def get_metrics() -> dict[str, int]:
    workspaces = list(store.workspaces.values())
    patches = [patch for workspace in workspaces for patch in workspace.patches]
    return {
        "schedules_imported": len(store.schedules),
        "reports_processed": len(workspaces),
        "events_extracted": sum(len(workspace.events) for workspace in workspaces),
        "updates_proposed": sum(len(patch.updates) for patch in patches),
        "patches_accepted": sum(patch.status == "ACCEPTED" for patch in patches),
    }


def run_dev_server(*, seed: bool = False) -> int:
    """Start the local server, optionally applying sample data first."""
    if seed:
        # Import only for explicit seed runs; app import and normal startup stay DB-free.
        from construction_reconciliation.database.seed import seed_database

        try:
            seed_database()
        except Exception as exc:
            # Driver error text can contain connection details; keep output generic.
            print(f"Database seeding failed ({type(exc).__name__}).", file=sys.stderr)
            return 1

    import uvicorn

    uvicorn.run(
        "construction_reconciliation.main:app",
        host="127.0.0.1",
        port=8000,
        reload=True,
    )
    return 0


def _main() -> int:
    parser = argparse.ArgumentParser(description="Run the local reconciliation API.")
    parser.add_argument(
        "--seed",
        action="store_true",
        help="apply the repeatable sample data before starting the API",
    )
    args = parser.parse_args()
    return run_dev_server(seed=args.seed)


if __name__ == "__main__":
    raise SystemExit(_main())
